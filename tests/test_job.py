import asyncio
from pathlib import Path

import pytest
from openpyxl import Workbook

from job import (
    JobBusyError,
    JobIdleError,
    JobManager,
    JobStartError,
    carrier_query_times,
    job_elapsed_ms,
)
from models import TrackResult


def _write_input(path: Path, pairs: list[tuple[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Container"
    sheet["B1"] = "Carrier"
    for index, (container, carrier) in enumerate(pairs, start=2):
        sheet.cell(index, 1, container)
        sheet.cell(index, 2, carrier)
    workbook.save(path)


@pytest.mark.asyncio
async def test_job_start_stop_and_busy(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU"), ("YMLU1234567", "YMJA")])
    started = asyncio.Event()

    async def fake_run_batch(rows, **kwargs):
        kwargs["on_progress"](
            {"index": 0, "total": len(rows), "phase": "querying", "result": None}
        )
        started.set()
        await kwargs["cancel_event"].wait()
        return [], kwargs["output_path"]

    manager = JobManager(
        run_batch_fn=fake_run_batch,
        input_path=source,
        output_path=tmp_path / "out.xlsx",
    )
    assert manager.snapshot()["job"]["total"] == 2

    manager.start()
    await asyncio.wait_for(started.wait(), timeout=2)
    snap = manager.snapshot()
    assert snap["job"]["state"] == "running"
    assert snap["counts"]["QUERYING"] == 1
    with pytest.raises(JobBusyError):
        manager.start()
    with pytest.raises(JobBusyError):
        manager.reload_from_disk()

    manager.request_stop()
    assert manager.snapshot()["job"]["state"] == "stopping"
    if manager.task is not None:
        await manager.task
    assert manager.snapshot()["job"]["state"] == "stopped"
    assert manager.snapshot()["job"]["message"].startswith("Stopped")


def test_job_start_rejects_empty(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [])
    manager = JobManager(
        run_batch_fn=lambda *a, **k: None,
        input_path=source,
        output_path=tmp_path / "out.xlsx",
    )
    with pytest.raises(JobStartError):
        manager.start()
    with pytest.raises(JobIdleError):
        manager.request_stop()


def test_job_start_rejects_empty_carrier_filter(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])
    manager = JobManager(
        run_batch_fn=lambda *a, **k: None,
        input_path=source,
        output_path=tmp_path / "out.xlsx",
    )
    with pytest.raises(JobStartError, match="No carrier selected"):
        manager.start(carriers=[])


def test_challenge_progress_exposes_action_and_clears_on_resume(tmp_path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("ECMU7271573", "CMDU")])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    manager._on_progress({
        "index": 0, "phase": "challenge",
        "challenge": {"code": "CAPTCHA", "mode": "current_browser", "timeout_seconds": 180},
    })
    snap = manager.snapshot()
    assert snap["rows"][0]["phase"] == "challenge"
    assert snap["counts"]["QUERYING"] == 1
    assert snap["counts"]["PENDING"] == 0
    assert "keep it open" in snap["job"]["message"]
    assert snap["job"]["challenge"]["code"] == "CAPTCHA"
    manager._on_progress({"index": 0, "phase": "querying"})
    assert manager.snapshot()["job"]["challenge"] is None


def test_challenge_progress_carrier_unlock_mentions_batch(tmp_path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("ECMU7271573", "CMDU"), ("CMAU7662786", "CMDU")])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    manager._on_progress({
        "index": 0, "phase": "challenge",
        "challenge": {
            "code": "CAPTCHA",
            "mode": "current_browser",
            "timeout_seconds": 600,
            "scope": "carrier",
        },
    })
    message = manager.snapshot()["job"]["message"]
    assert "Batch query" in message
    assert "keep it open" in message
    assert "up to 10 min" in message


def test_parallel_progress_keeps_all_active_rows(tmp_path):
    source = tmp_path / "in.xlsx"
    _write_input(
        source,
        [("HLXU1234567", "HLCU"), ("YMLU1234567", "YMJA")],
    )
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    manager._on_progress(
        {"index": 0, "phase": "querying", "challenge": None}
    )
    manager._on_progress(
        {
            "index": 1,
            "phase": "challenge",
            "challenge": {
                "code": "CAPTCHA",
                "mode": "automatic",
                "timeout_seconds": 25,
            },
        }
    )

    snap = manager.snapshot()
    assert snap["counts"]["QUERYING"] == 2
    assert snap["job"]["current_indices"] == [0, 1]
    assert snap["job"]["challenges"][0]["index"] == 1
    assert [row["phase"] for row in snap["rows"]] == ["querying", "challenge"]

    result = TrackResult(
        container="HLXU1234567",
        carrier="HLCU",
        status="NOT_LOADED",
        checked_at="t",
    )
    manager._on_progress(
        {"index": 0, "phase": "done", "result": result}
    )
    snap = manager.snapshot()
    assert snap["counts"]["QUERYING"] == 1
    assert snap["job"]["current_indices"] == [1]


def test_job_snapshot_merges_previous_results(tmp_path: Path):
    from excel_io import build_output_frame, write_output

    source = tmp_path / "in.xlsx"
    output = tmp_path / "out.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])
    write_output(
        output,
        build_output_frame(
            [{"Container": "HLXU1234567", "Carrier": "HLCU", "extras": {}}],
            [
                TrackResult(
                    container="HLXU1234567",
                    carrier="HLCU",
                    status="SAILED",
                    success=True,
                    check_result="SUCCESS",
                    loaded=True,
                    sailed=True,
                    checked_at="2026-09-13 00:00:00",
                )
            ],
        ),
    )
    manager = JobManager(input_path=source, output_path=output)
    row = manager.snapshot()["rows"][0]
    assert row["status"] == "SAILED"
    assert row["phase"] == "done"
    assert manager.snapshot()["carrier_times"] == {}


def test_job_elapsed_ms_from_start_until_now_or_finish():
    assert job_elapsed_ms(None, None, now_ms=5_000) is None
    assert job_elapsed_ms(1_000, None, now_ms=5_000) == 4_000
    assert job_elapsed_ms(1_000, 3_000, now_ms=9_000) == 2_000
    assert job_elapsed_ms(5_000, 4_000, now_ms=9_000) == 0


def test_job_snapshot_exposes_elapsed_ms(tmp_path: Path, monkeypatch):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])
    clock = {"t": 1_000_000}
    monkeypatch.setattr("job._now_ms", lambda: clock["t"])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    assert manager.snapshot()["job"]["elapsed_ms"] is None

    manager.started_ms = 1_000_000
    clock["t"] = 1_012_500
    assert manager.snapshot()["job"]["elapsed_ms"] == 12_500
    manager.finished_ms = 1_020_000
    assert manager.snapshot()["job"]["elapsed_ms"] == 20_000


def test_carrier_query_times_use_wall_clock_and_average():
    rows = [
        {"Carrier": "HLCU"},
        {"Carrier": "HLCU"},
        {"Carrier": "YMJA"},
    ]
    times = {
        0: {"started_ms": 1_000, "finished_ms": 11_000},
        1: {"started_ms": 13_000, "finished_ms": 21_000},
        2: {"started_ms": 2_000, "finished_ms": 6_000},
    }
    assert carrier_query_times(rows, times, now_ms=30_000) == {
        "HLCU": {"total_ms": 20_000, "avg_ms": 9_000, "queried": 2},
        "YMJA": {"total_ms": 4_000, "avg_ms": 4_000, "queried": 1},
    }


def test_job_snapshot_records_query_times(tmp_path: Path, monkeypatch):
    source = tmp_path / "in.xlsx"
    _write_input(
        source,
        [("HLXU1234567", "HLCU"), ("HLXU7654321", "HLCU")],
    )
    clock = {"t": 1_000_000}
    monkeypatch.setattr("job._now_ms", lambda: clock["t"])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    assert manager.snapshot()["carrier_times"] == {}

    manager._on_progress({"index": 0, "phase": "querying"})
    clock["t"] = 1_010_000
    manager._on_progress(
        {
            "index": 0,
            "phase": "done",
            "result": TrackResult(
                container="HLXU1234567",
                carrier="HLCU",
                status="NOT_LOADED",
                checked_at="t",
            ),
        }
    )
    clock["t"] = 1_012_000
    manager._on_progress({"index": 1, "phase": "querying"})
    clock["t"] = 1_020_000
    snap = manager.snapshot()
    assert snap["carrier_times"]["HLCU"] == {
        "total_ms": 20_000,
        "avg_ms": 9_000,
        "queried": 2,
    }

    manager.reload_from_disk()
    assert manager.snapshot()["carrier_times"] == {}


def test_done_without_query_is_not_timed(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    manager._on_progress(
        {
            "index": 0,
            "phase": "done",
            "result": TrackResult(
                container="HLXU1234567",
                carrier="HLCU",
                status="SAILED",
                checked_at="t",
            ),
        }
    )
    assert manager.snapshot()["carrier_times"] == {}
