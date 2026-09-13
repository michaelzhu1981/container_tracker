from pathlib import Path

from openpyxl import load_workbook

from config import INPUT_XLSX
from excel_io import (
    build_output_frame,
    create_input_template,
    order_rows_by_carrier,
    read_input,
    write_output,
)
from models import TrackResult


def test_read_and_write_roundtrip(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "箱号"
    ws["B1"] = "Carrier"
    ws["C1"] = "Customer"
    ws["A2"] = "hlxu-1234567"
    ws["B2"] = "hlcu"
    ws["C2"] = "ACME"
    wb.save(source)

    rows = read_input(source)
    assert len(rows) == 1
    assert rows[0]["Container"] == "HLXU1234567"
    assert rows[0]["Carrier"] == "HLCU"
    assert rows[0]["extras"]["Customer"] == "ACME"

    result = TrackResult(
        container="HLXU1234567",
        carrier="HLCU",
        pol="YANTIAN",
        loaded=True,
        sailed=True,
        status="SAILED",
        success=True,
        check_result="SUCCESS",
        checked_at="2026-09-12 00:00:00",
    )
    frame = build_output_frame(rows, [result])
    out = write_output(tmp_path / "out.xlsx", frame)
    assert out.exists()
    again = read_input(out)
    assert again[0]["Container"] == "HLXU1234567"


def test_read_input_dedupes_container_carrier_by_default(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Container"
    ws["B1"] = "Carrier"
    ws["A2"] = "CAIU7012411"
    ws["B2"] = "HLCU"
    ws["A3"] = "BMOU5733569"
    ws["B3"] = "YMJA"
    ws["A4"] = "CAIU7012411"
    ws["B4"] = "HLCU"
    wb.save(source)

    rows = read_input(source)
    assert [(r["Container"], r["Carrier"]) for r in rows] == [
        ("CAIU7012411", "HLCU"),
        ("BMOU5733569", "YMJA"),
    ]


def test_order_rows_by_carrier_keeps_first_seen_groups():
    rows = [
        {"Container": "A", "Carrier": "YMJA"},
        {"Container": "B", "Carrier": "HLCU"},
        {"Container": "C", "Carrier": "YMJA"},
        {"Container": "D", "Carrier": "HLCU"},
    ]
    ordered = order_rows_by_carrier(rows)
    assert [r["Container"] for r in ordered] == ["A", "C", "B", "D"]


def test_create_input_template(tmp_path: Path):
    path = create_input_template(tmp_path / "containers.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Containers", "Instructions"]
    sheet = workbook["Containers"]
    assert sheet["A1"].value == "Container"
    assert sheet["B1"].value == "Carrier"
    assert sheet["A1"].font.bold is True
    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value in (None, "")
    assert not sheet.data_validations.dataValidation
    rows = read_input(path)
    assert rows == []
    instructions = workbook["Instructions"]
    assert "python app.py input/containers.xlsx" in str(instructions["A3"].value)
    assert "OOLU, HDMU, COSU, or ZIMU" in str(instructions["A2"].value)


def test_repo_input_template_exists():
    assert INPUT_XLSX.exists()
    workbook = load_workbook(INPUT_XLSX)
    assert workbook.sheetnames[0] == "Containers"
