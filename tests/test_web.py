from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from job import JobManager
from web import create_app


def _write_input(path: Path, pairs: list[tuple[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Container"
    sheet["B1"] = "Carrier"
    for index, (container, carrier) in enumerate(pairs, start=2):
        sheet.cell(index, 1, container)
        sheet.cell(index, 2, carrier)
    workbook.save(path)


def test_status_page_and_stop_when_idle(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])
    manager = JobManager(input_path=source, output_path=tmp_path / "out.xlsx")
    with TestClient(create_app(manager)) as client:
        page = client.get("/")
        assert page.status_code == 200
        assert "Container Tracker" in page.text
        assert ">Time<" in page.text
        status = client.get("/api/status")
        assert status.status_code == 200
        assert status.json()["rows"][0]["container"] == "HLXU1234567"
        assert status.json()["carrier_times"] == {}
        stopped = client.post("/api/stop")
        assert stopped.status_code == 400


def test_start_and_stop_job(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(source, [("HLXU1234567", "HLCU")])

    async def fake_run_batch(rows, **kwargs):
        kwargs["on_progress"](
            {"index": 0, "total": len(rows), "phase": "querying", "result": None}
        )
        await kwargs["cancel_event"].wait()
        return [], kwargs["output_path"]

    manager = JobManager(
        run_batch_fn=fake_run_batch,
        input_path=source,
        output_path=tmp_path / "out.xlsx",
    )
    with TestClient(create_app(manager)) as client:
        begun = client.post("/api/start", json={"resume": False})
        assert begun.status_code == 200
        assert begun.json()["job"]["state"] == "running"
        assert client.post("/api/start", json={}).status_code == 409
        assert client.post("/api/reload").status_code == 409
        for _ in range(50):
            if client.get("/api/status").json()["counts"]["QUERYING"] == 1:
                break
        halted = client.post("/api/stop")
        assert halted.status_code == 200
        assert halted.json()["job"]["state"] == "stopping"
        for _ in range(50):
            if client.get("/api/status").json()["job"]["state"] == "stopped":
                break
        assert client.get("/api/status").json()["job"]["state"] == "stopped"


def test_status_exposes_multiple_parallel_queries(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    _write_input(
        source,
        [("HLXU1234567", "HLCU"), ("YMLU1234567", "YMJA")],
    )

    async def fake_run_batch(rows, **kwargs):
        for index in range(2):
            kwargs["on_progress"](
                {
                    "index": index,
                    "total": len(rows),
                    "phase": "querying",
                    "result": None,
                }
            )
        await kwargs["cancel_event"].wait()
        return [], kwargs["output_path"]

    manager = JobManager(
        run_batch_fn=fake_run_batch,
        input_path=source,
        output_path=tmp_path / "out.xlsx",
    )
    with TestClient(create_app(manager)) as client:
        assert client.post("/api/start", json={}).status_code == 200
        status = {}
        for _ in range(50):
            status = client.get("/api/status").json()
            if status["counts"]["QUERYING"] == 2:
                break
        assert status["job"]["current_indices"] == [0, 1]
        assert status["counts"]["QUERYING"] == 2
        assert client.post("/api/stop").status_code == 200
