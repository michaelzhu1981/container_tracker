"""Excel read/write with copy-on-read and lock-safe replace."""

from __future__ import annotations

import shutil
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

from artifacts import relative_to_root
from config import HEADER_ALIASES, INPUT_XLSX, RESULT_COLUMNS, STANDARD_OUTPUT_NAMES
from models import TrackResult
from validate import (
    carrier_supported,
    container_shape_ok,
    iso6346_check_digit_ok,
    normalize_carrier,
    normalize_container,
)

YES_NO = {True: "YES", False: "NO", None: ""}


class ExcelReadError(Exception):
    pass


def _canonical_header(name: str) -> str:
    raw = str(name).strip()
    mapped = HEADER_ALIASES.get(raw.lower(), HEADER_ALIASES.get(raw, raw))
    return mapped


def read_input(path: Path) -> list[dict]:
    if path.suffix.lower() != ".xlsx":
        raise ExcelReadError("Input must be an .xlsx file.")
    if not path.exists():
        raise ExcelReadError(f"Input file not found: {path}")

    tmp = Path(tempfile.mkstemp(suffix=".xlsx")[1])
    try:
        shutil.copy2(path, tmp)
    except OSError as exc:
        raise ExcelReadError(
            "Cannot read the input file. Close it in Excel if copy failed, then retry."
        ) from exc

    try:
        frame = pd.read_excel(tmp, sheet_name=0, dtype=str, engine="openpyxl")
    except Exception as exc:  # noqa: BLE001
        raise ExcelReadError(f"Cannot parse input Excel: {exc}") from exc
    finally:
        tmp.unlink(missing_ok=True)

    frame = frame.rename(columns={col: _canonical_header(col) for col in frame.columns})
    if "Container" not in frame.columns or "Carrier" not in frame.columns:
        raise ExcelReadError("Missing required headers: Container and Carrier.")

    extra = [c for c in frame.columns if c not in {"Container", "Carrier"}]
    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for _, raw in frame.iterrows():
        container = normalize_container(raw.get("Container", "") or "")
        carrier = normalize_carrier(raw.get("Carrier", "") or "")
        if not container and not carrier:
            continue
        extras = {}
        for col in extra:
            key = f"{col}_input" if col in STANDARD_OUTPUT_NAMES or col == "POL" else col
            if col == "POL":
                key = "POL_input"
            extras[key] = "" if pd.isna(raw.get(col)) else str(raw.get(col)).strip()
        item = {
            "Container": container,
            "Carrier": carrier,
            "extras": extras,
            "shape_ok": container_shape_ok(container),
            "checksum_ok": iso6346_check_digit_ok(container) if container_shape_ok(container) else False,
            "carrier_ok": carrier_supported(carrier),
        }
        key = (container, carrier)
        if key in seen:
            continue
        seen.add(key)
        rows.append(item)
    return rows


def order_rows_by_carrier(rows: list[dict]) -> list[dict]:
    """Keep first-seen carrier order; consecutive boxes of the same carrier."""
    by_carrier: dict[str, list[dict]] = {}
    for row in rows:
        by_carrier.setdefault(row["Carrier"], []).append(row)
    return [row for group in by_carrier.values() for row in group]


def result_to_cells(result: TrackResult | None) -> dict[str, str]:
    empty = {col: "" for col in RESULT_COLUMNS}
    if result is None:
        return empty
    empty.update(
        {
            "Container": result.container,
            "Carrier": result.carrier,
            "POL": result.pol or "",
            "Status": result.status or "",
            "Loaded": YES_NO[result.loaded],
            "Sailed": YES_NO[result.sailed],
            "Vessel": result.vessel or "",
            "Voyage": result.voyage or "",
            "ATD": result.atd or "",
            "Latest Event": result.latest_event or "",
            "Checked At": result.checked_at or "",
            "Check Result": result.check_result or "",
            "Error Code": result.error_code or "",
            "Error": result.error or "",
            "Screenshot": result.screenshot_path or "",
        }
    )
    return empty


def build_output_frame(input_rows: list[dict], results: list[TrackResult | None]) -> pd.DataFrame:
    records = []
    extras_keys: list[str] = []
    for row, result in zip(input_rows, results, strict=True):
        cells = result_to_cells(result)
        cells["Container"] = row["Container"]
        cells["Carrier"] = row["Carrier"]
        extras = row.get("extras") or {}
        for key in extras:
            if key not in extras_keys:
                extras_keys.append(key)
            cells[key] = extras[key]
        records.append(cells)
    columns = list(RESULT_COLUMNS) + extras_keys
    if not records:
        return pd.DataFrame(columns=columns)
    frame = pd.DataFrame(records)
    return frame.reindex(columns=columns)


def write_output(path: Path, frame: pd.DataFrame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    partial = path.with_suffix(path.suffix + ".partial")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    header_font = Font(bold=True)
    for r_idx, row in enumerate(dataframe_to_rows(frame.fillna(""), index=False, header=True), start=1):
        for c_idx, value in enumerate(row, start=1):
            cell = sheet.cell(r_idx, c_idx, "" if value is None else str(value))
            cell.number_format = "@"
            if r_idx == 1:
                cell.font = header_font
    workbook.save(partial)
    try:
        partial.replace(path)
        return path
    except OSError:
        stamped = path.with_name(
            f"{path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{path.suffix}"
        )
        try:
            partial.replace(stamped)
        except OSError:
            shutil.copy2(partial, stamped)
            partial.unlink(missing_ok=True)
        shown = relative_to_root(stamped) or str(stamped)
        print(f"Output file is open in Excel. Wrote {shown} instead.")
        return stamped


def create_input_template(path: Path | None = None) -> Path:
    """Create the empty Containers + Instructions workbook used as the input template."""
    target = path or INPUT_XLSX
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Containers"
    sheet["A1"] = "Container"
    sheet["B1"] = "Carrier"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"].font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 12
    for row in range(2, 102):
        for col in (1, 2):
            cell = sheet.cell(row, col, "")
            cell.number_format = "@"
    instructions = workbook.create_sheet("Instructions")
    lines = (
        "Paste container numbers into column A starting at row 2.",
        "Type a Carrier code in column B for each row: HLCU, YMJA, ONEY, MAEU, MSCU, or CMDU.",
        "Save this file, then run: python app.py input/containers.xlsx",
        "Do not fill POL. The program infers POL from the latest ocean journey.",
    )
    for index, line in enumerate(lines, start=1):
        instructions.cell(index, 1, line)
    instructions.column_dimensions["A"].width = 100
    workbook.save(target)
    return target


def load_previous_results(path: Path) -> dict[tuple[str, str, int], dict]:
    if not path.exists():
        return {}
    tmp = Path(tempfile.mkstemp(suffix=".xlsx")[1])
    try:
        shutil.copy2(path, tmp)
        frame = pd.read_excel(tmp, sheet_name=0, dtype=str, engine="openpyxl")
    except OSError:
        return {}
    finally:
        tmp.unlink(missing_ok=True)
    previous: dict[tuple[str, str, int], dict] = {}
    counts: dict[tuple[str, str], int] = {}
    for _, raw in frame.iterrows():
        container = normalize_container(raw.get("Container", "") or "")
        carrier = normalize_carrier(raw.get("Carrier", "") or "")
        occurrence = counts.get((container, carrier), 0)
        counts[(container, carrier)] = occurrence + 1
        previous[(container, carrier, occurrence)] = {
            col: "" if pd.isna(raw.get(col, "")) else str(raw.get(col, "")).strip()
            for col in frame.columns
        }
    return previous
